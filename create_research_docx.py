# ============================================================
# EDUPRO - PROFESSIONAL RESEARCH PAPER DOCX GENERATOR
# STEP 32
# ============================================================

import os
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# PATHS
# ============================================================

REPORT_DIR = os.path.join(
    "outputs",
    "reports"
)

TXT_FILE = os.path.join(
    REPORT_DIR,
    "EduPro_Final_Research_Paper.txt"
)

DOCX_FILE = os.path.join(
    REPORT_DIR,
    "EduPro_Final_Research_Paper.docx"
)


# ============================================================
# CHECK INPUT
# ============================================================

if not os.path.exists(TXT_FILE):

    print()
    print("ERROR: Research paper text file not found.")
    print()
    print(
        "First run:"
    )
    print(
        "python research_paper.py"
    )
    print()

    raise SystemExit(1)


# ============================================================
# READ PAPER
# ============================================================

with open(
    TXT_FILE,
    "r",
    encoding="utf-8"
) as file:

    paper_text = file.read()


# ============================================================
# CREATE DOCUMENT
# ============================================================

doc = Document()


# ============================================================
# PAGE SETTINGS
# ============================================================

for section in doc.sections:

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)


# ============================================================
# DEFAULT FONT
# ============================================================

normal_style = doc.styles["Normal"]

normal_style.font.name = "Times New Roman"
normal_style.font.size = Pt(11)

normal_style.paragraph_format.line_spacing = 1.15
normal_style.paragraph_format.space_after = Pt(6)


# ============================================================
# HEADING STYLES
# ============================================================

heading1 = doc.styles["Heading 1"]

heading1.font.name = "Times New Roman"
heading1.font.size = Pt(16)
heading1.font.bold = True

heading1.paragraph_format.space_before = Pt(14)
heading1.paragraph_format.space_after = Pt(8)


heading2 = doc.styles["Heading 2"]

heading2.font.name = "Times New Roman"
heading2.font.size = Pt(13)
heading2.font.bold = True


# ============================================================
# HEADER
# ============================================================

section = doc.sections[0]

header = section.header.paragraphs[0]

header.alignment = WD_ALIGN_PARAGRAPH.CENTER

header_run = header.add_run(
    "EduPro Learner Analytics — Research Project"
)

header_run.font.name = "Times New Roman"
header_run.font.size = Pt(9)
header_run.italic = True


# ============================================================
# FOOTER WITH PAGE NUMBER
# ============================================================

footer = section.footer.paragraphs[0]

footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = footer.add_run(
    "Page "
)

run.font.name = "Times New Roman"
run.font.size = Pt(9)

field = OxmlElement(
    "w:fldSimple"
)

field.set(
    qn("w:instr"),
    "PAGE"
)

footer._p.append(field)


# ============================================================
# COVER PAGE
# ============================================================

p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.paragraph_format.space_before = Pt(70)

r = p.add_run(
    "RESEARCH PROJECT REPORT"
)

r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(18)


p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.paragraph_format.space_before = Pt(30)

r = p.add_run(
    "LEARNER DEMOGRAPHICS AND COURSE\n"
    "ENROLLMENT BEHAVIOR ANALYSIS ON EDUPRO"
)

r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(22)


p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.paragraph_format.space_before = Pt(20)

r = p.add_run(
    "A Descriptive Analytics Approach for Understanding\n"
    "Learner Participation, Course Preferences and\n"
    "Enrollment Behavior"
)

r.italic = True
r.font.name = "Times New Roman"
r.font.size = Pt(13)


p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.paragraph_format.space_before = Pt(60)

r = p.add_run(
    "Master of Computer Applications (MCA)\n\n"
    "Research Project"
)

r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(13)


p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.paragraph_format.space_before = Pt(40)

r = p.add_run(
    "Prepared for Academic Presentation and Evaluation"
)

r.font.name = "Times New Roman"
r.font.size = Pt(11)


# ============================================================
# PAGE BREAK
# ============================================================

doc.add_page_break()


# ============================================================
# TABLE OF CONTENTS
# ============================================================

doc.add_heading(
    "Table of Contents",
    level=1
)

toc_items = [

    "1. Abstract",

    "2. Keywords",

    "3. Introduction",

    "4. Problem Statement",

    "5. Research Objectives",

    "6. Research Questions",

    "7. Dataset Description",

    "8. Methodology",

    "9. Data Preprocessing",

    "10. Descriptive Analysis",

    "11. Course Preference Analysis",

    "12. Temporal Analysis",

    "13. Cross-Dimensional Analysis",

    "14. Interactive Dashboard",

    "15. Results and Findings",

    "16. Recommendations",

    "17. SDG 4 Alignment",

    "18. Government Stakeholder Relevance",

    "19. Limitations",

    "20. Future Scope",

    "21. Conclusion",

    "22. Final Research Statement",

    "References"
]


for item in toc_items:

    p = doc.add_paragraph()

    p.paragraph_format.left_indent = Inches(0.2)

    p.add_run(
        item
    )


doc.add_page_break()


# ============================================================
# PARSE PAPER CONTENT
# ============================================================

lines = paper_text.splitlines()


for line in lines:

    text = line.strip()

    # Skip empty lines
    if not text:
        continue

    # Skip separator lines
    if set(text) == {"="}:
        continue

    if set(text) == {"-"}:
        continue

    # Skip original title/cover information
    if text.startswith(
        "LEARNER DEMOGRAPHICS AND COURSE ENROLLMENT"
    ):
        continue

    if text.startswith(
        "BEHAVIOR ANALYSIS ON EDUPRO"
    ):
        continue

    if text.startswith(
        "A DESCRIPTIVE ANALYTICS APPROACH"
    ):
        continue

    if text.startswith(
        "Research Project –"
    ):
        continue

    # --------------------------------------------------------
    # Numbered section
    # --------------------------------------------------------

    section_match = re.match(
        r"^(\d+)\.\s+(.+)$",
        text
    )

    if section_match:

        number = section_match.group(1)
        title = section_match.group(2)

        p = doc.add_heading(
            f"{number}. {title}",
            level=1
        )

        continue

    # --------------------------------------------------------
    # Research Question
    # --------------------------------------------------------

    if re.match(
        r"^RQ\d+\.",
        text
    ):

        p = doc.add_paragraph()

        p.paragraph_format.left_indent = Inches(0.2)

        run = p.add_run(
            text
        )

        run.bold = True

        continue

    # --------------------------------------------------------
    # Bullet
    # --------------------------------------------------------

    if text.startswith("•"):

        p = doc.add_paragraph(
            text[1:].strip(),
            style="List Bullet"
        )

        continue

    # --------------------------------------------------------
    # Numbered list
    # --------------------------------------------------------

    numbered = re.match(
        r"^(\d+)\.\s+(.+)$",
        text
    )

    if numbered:

        p = doc.add_paragraph(
            numbered.group(2),
            style="List Number"
        )

        continue

    # --------------------------------------------------------
    # Label / Subheading
    # --------------------------------------------------------

    if text.endswith(":") and len(text) < 100:

        p = doc.add_paragraph()

        run = p.add_run(
            text
        )

        run.bold = True

        continue

    # --------------------------------------------------------
    # Normal paragraph
    # --------------------------------------------------------

    p = doc.add_paragraph()

    p.paragraph_format.first_line_indent = Inches(
        0.25
    )

    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p.add_run(
        text
    )


# ============================================================
# REFERENCES
# ============================================================

doc.add_page_break()

doc.add_heading(
    "References",
    level=1
)

references = [

    "United Nations. Sustainable Development Goals — Goal 4: Quality Education.",

    "Streamlit Documentation. Streamlit Application Development and API Documentation.",

    "Python Software Foundation. Python Documentation.",

    "Pandas Development Team. Pandas Documentation.",

    "Plotly Technologies Inc. Plotly Python Graphing Library Documentation.",

    "OpenPyXL Documentation. Python Library for Reading and Writing Excel Files."
]


for reference in references:

    p = doc.add_paragraph(
        reference,
        style="List Number"
    )

    p.paragraph_format.space_after = Pt(6)


# ============================================================
# SAVE DOCUMENT
# ============================================================

os.makedirs(
    REPORT_DIR,
    exist_ok=True
)

doc.save(
    DOCX_FILE
)


# ============================================================
# FINAL OUTPUT
# ============================================================

print()
print("=" * 70)
print("STEP 32 COMPLETED SUCCESSFULLY")
print("=" * 70)
print()

print(
    "Professional research paper created:"
)

print(
    DOCX_FILE
)

print()

print(
    "Document includes:"
)

print(
    "✓ Professional Cover Page"
)

print(
    "✓ Table of Contents"
)

print(
    "✓ Abstract"
)

print(
    "✓ Keywords"
)

print(
    "✓ Introduction"
)

print(
    "✓ Problem Statement"
)

print(
    "✓ Research Objectives"
)

print(
    "✓ Research Questions"
)

print(
    "✓ Dataset Description"
)

print(
    "✓ Methodology"
)

print(
    "✓ Data Preprocessing"
)

print(
    "✓ Descriptive Analysis"
)

print(
    "✓ Course Preference Analysis"
)

print(
    "✓ Temporal Analysis"
)

print(
    "✓ Cross-Dimensional Analysis"
)

print(
    "✓ Dashboard Description"
)

print(
    "✓ Results and Findings"
)

print(
    "✓ Recommendations"
)

print(
    "✓ SDG 4 Alignment"
)

print(
    "✓ Government Stakeholder Relevance"
)

print(
    "✓ Limitations"
)

print(
    "✓ Future Scope"
)

print(
    "✓ Conclusion"
)

print(
    "✓ References"
)

print()
print("=" * 70)